import os
import copy
import json
import warnings
import logging
import numpy as np
import pandas as pd
from datetime import datetime, timedelta
import time
import matplotlib.pyplot as plt
import pywt
from scipy.signal import medfilt
from scipy.stats import skew, kurtosis
from sklearn.preprocessing import LabelEncoder, StandardScaler, PowerTransformer
from sklearn.cluster import KMeans
from sklearn.neighbors import LocalOutlierFactor
import xgboost as xgb
import lightgbm as lgb

warnings.filterwarnings("ignore")
logging.getLogger("xgboost").setLevel(logging.ERROR)
logging.getLogger("lightgbm").setLevel(logging.ERROR)


def mae(y_true, y_pred):
    return np.mean(np.abs(y_true - y_pred))


def smape(y_true, y_pred):
    return 100 * np.mean(
        np.abs(y_pred - y_true) /
        (np.abs(y_pred) + np.abs(y_true) + 1e-6)
    )


def make_safe_name(text):
    text = str(text)
    for old, new in [
        ("=", "_eq_"), (".", "p"), (" ", "_"),
        ("/", "_"), ("\\", "_"), (":", "_"),
        ("[", "_"), ("]", "_"), ("(", "_"), (")", "_"),
        (",", "_"), (";", "_"), ("'", "_"), ('"', "_")
    ]:
        text = text.replace(old, new)
    return text


def trimmed_stats(residuals, weights, trim_ratio=0.05):
    r = np.array(residuals, dtype=float)
    w = np.array(weights, dtype=float)

    if len(r) == 0:
        return r, w

    n = len(r)
    k = int(n * trim_ratio)
    if k < 1 or 2 * k >= n:
        return r, w

    idx = np.argsort(np.abs(r))
    keep_idx = idx[k: n - k]
    return r[keep_idx], w[keep_idx]


def compute_huber_alpha(residuals, scale=2.0):
    med = np.median(np.abs(residuals)) + 1e-6
    sigma = med * 1.4826 * scale
    return float(max(sigma, 1e-3))


def compute_dynamic_max_depth_advanced(
        residuals, weights,
        d0=6,
        delta_pos=1.5,
        delta_neg=1.2,
        skew_coef=5.0,
        kurt_coef=5.0,
        min_depth=2,
        max_depth_limit=20,
        prev_depth=None,
        trim_ratio=0.05,
        smooth_factor=0.3,
        return_detail=False):
    r_trim, w_trim = trimmed_stats(residuals, weights, trim_ratio=trim_ratio)
    if len(r_trim) == 0:
        if return_detail:
            return d0, {
                "sigma_w": 0.0, "skew_val": 0.0, "kurt_val": 0.0,
                "mad_w": 0.0, "f": 0.0, "raw": float(d0)
            }
        return d0

    w_sum = np.sum(w_trim) + 1e-12
    mu = np.sum(w_trim * r_trim) / w_sum
    var_w = np.sum(w_trim * (r_trim - mu) ** 2) / w_sum
    sigma_w = np.sqrt(var_w)

    skew_val = abs(skew(r_trim)) if len(r_trim) > 2 else 0.0
    kurt_val = abs(kurtosis(r_trim, fisher=True)) if len(r_trim) > 3 else 0.0

    med_abs = np.median(np.abs(r_trim - np.median(r_trim))) + 1e-6
    mad_w = med_abs * 1.4826

    f = (sigma_w * (1.0 + skew_coef * skew_val + kurt_coef * kurt_val)) ** (1 / 3) + (mad_w) ** (1 / 3)

    if f >= 1.0:
        raw = d0 + delta_pos * f
    else:
        raw = d0 - delta_neg * f

    if prev_depth is not None:
        raw = smooth_factor * raw + (1.0 - smooth_factor) * prev_depth

    depth = int(round(raw))
    depth = max(min_depth, min(depth, max_depth_limit))

    if return_detail:
        return depth, {
            "sigma_w": float(sigma_w),
            "skew_val": float(skew_val),
            "kurt_val": float(kurt_val),
            "mad_w": float(mad_w),
            "f": float(f),
            "raw": float(raw)
        }

    return depth


def compute_dynamic_gamma_advanced(
        residuals, weights,
        base=1.0,
        gamma_alpha=0.7,
        skew_coef=5.0,
        kurt_coef=5.0,
        min_gamma=0.001,
        max_gamma=20.0,
        prev_gamma=None,
        trim_ratio=0.05,
        smooth_factor=0.3,
        return_detail=False):
    r_arr = np.array(residuals, dtype=float)
    w_arr = np.array(weights, dtype=float)

    r_trim, w_trim = trimmed_stats(r_arr, w_arr, trim_ratio=trim_ratio)
    if len(r_trim) == 0:
        if return_detail:
            return base, {
                "iqr_w": 0.0, "mad_w": 0.0, "skew_val": 0.0,
                "kurt_val": 0.0, "f_gamma": 0.0, "raw": float(base)
            }
        return base

    idx = np.argsort(r_trim)
    r_sorted = r_trim[idx]
    w_sorted = w_trim[idx]
    w_cum = np.cumsum(w_sorted)
    w_sum = w_cum[-1] + 1e-12

    q1_pos = 0.25 * w_sum
    q3_pos = 0.75 * w_sum
    q1_idx = np.searchsorted(w_cum, q1_pos)
    q3_idx = np.searchsorted(w_cum, q3_pos)

    r_q1 = r_sorted[min(q1_idx, len(r_sorted) - 1)]
    r_q3 = r_sorted[min(q3_idx, len(r_sorted) - 1)]
    iqr_w = abs(r_q3 - r_q1)

    med_abs = np.median(np.abs(r_trim - np.median(r_trim))) + 1e-6
    mad_w = med_abs * 1.4826

    skew_val = abs(skew(r_trim)) if len(r_trim) > 2 else 0.0
    kurt_val = abs(kurtosis(r_trim, fisher=True)) if len(r_trim) > 3 else 0.0

    f_gamma = (iqr_w / (mad_w + 1e-12)) * (1.0 + skew_coef * skew_val + kurt_coef * kurt_val)

    raw = base / (1.0 + gamma_alpha * f_gamma)

    if prev_gamma is not None:
        raw = smooth_factor * raw + (1.0 - smooth_factor) * prev_gamma

    gamma = float(max(min_gamma, min(raw, max_gamma)))

    if return_detail:
        return gamma, {
            "iqr_w": float(iqr_w),
            "mad_w": float(mad_w),
            "skew_val": float(skew_val),
            "kurt_val": float(kurt_val),
            "f_gamma": float(f_gamma),
            "raw": float(raw)
        }

    return gamma


def wavelet_denoise_series(arr, wavelet="db4", level=3, thr=None):
    arr = np.asarray(arr, dtype=float)
    if len(arr) < 8:
        return arr.copy(), (0.0 if thr is None else float(thr))

    wavelet_obj = pywt.Wavelet(wavelet)
    max_level = pywt.dwt_max_level(len(arr), wavelet_obj.dec_len)
    use_level = min(level, max_level) if max_level >= 1 else 1

    coeffs = pywt.wavedec(arr, wavelet, level=use_level)

    if thr is None:
        sigma = np.median(np.abs(coeffs[-1])) / 0.6745 + 1e-12
        thr = sigma * np.sqrt(2 * np.log(len(arr)))

    coeffs[1:] = [pywt.threshold(c, thr, 'soft') for c in coeffs[1:]]
    rec = pywt.waverec(coeffs, wavelet)[:len(arr)]
    return rec, float(thr)


def save_convergence_curve(values, title, save_path, best_round=None):
    fig, ax = plt.subplots(figsize=(6, 4))
    x = np.arange(1, len(values) + 1)
    ax.plot(x, values, marker='o', linewidth=1)
    if best_round is not None and 1 <= best_round <= len(values):
        ax.axvline(best_round, linestyle='--', linewidth=1)
    ax.set_title(title)
    ax.set_xlabel("Round")
    ax.set_ylabel("Validation MAE")
    plt.tight_layout()
    plt.savefig(save_path)
    plt.close(fig)


def save_sensitivity_plots(summary_df, plot_dir):
    os.makedirs(plot_dir, exist_ok=True)

    for param_name in summary_df["param_name"].dropna().unique():
        sub = summary_df[summary_df["param_name"] == param_name].copy()
        if sub.empty:
            continue

        sub["param_value_num"] = pd.to_numeric(sub["param_value"], errors="coerce")
        if sub["param_value_num"].notna().all():
            sub = sub.sort_values("param_value_num")
            x = sub["param_value_num"].values
            xlabel = param_name
        else:
            sub = sub.reset_index(drop=True)
            x = np.arange(len(sub))
            xlabel = param_name

        plt.figure(figsize=(7, 4))
        plt.plot(x, sub["mae_pt_mean"].values, marker="o", label="parking_time MAE")
        plt.plot(x, sub["mae_kWh_mean"].values, marker="s", label="kWhDelivered MAE")
        if not sub["param_value_num"].notna().all():
            plt.xticks(x, sub["param_value"].astype(str), rotation=45)
        plt.xlabel(xlabel)
        plt.ylabel("Mean MAE")
        plt.title(f"Sensitivity of {param_name}")
        plt.legend()
        plt.tight_layout()
        plt.savefig(os.path.join(plot_dir, f"{make_safe_name(param_name)}_sensitivity.png"))
        plt.close()

def get_lgb_params(random_state=42):
    return {
        "objective": "regression",
        "metric": "mae",
        "verbosity": -1,
        "boosting_type": "gbdt",
        "n_estimators": 200,
        "learning_rate": 0.02,
        "num_leaves": 31,
        "feature_fraction": 0.8,
        "bagging_fraction": 0.8,
        "bagging_freq": 1,
        "min_data_in_leaf": 20,
        "lambda_l1": 0.1,
        "lambda_l2": 0.1,
        "random_state": random_state
    }


def get_base_config():
    return {
        "CSV_PATH": "caltech_test_data.csv",
        "RESULTS_ROOT": "./results/caltech_joint_sensitivity",

        "RUN_MODE": "confirm",

        "ENABLE_SENSITIVITY": True,

        "WINDOWS": [30, 60, 120, 240, 360, 480],

        "BASE_SEED": 2026,
        "EPS": 1e-3,
        "TEST_START": "2019-12-01",
        "TEST_END": "2019-12-30",

        "PT_FIXED_DEPTH": 6,
        "PT_FIXED_GAMMA": 1.0,
        "PT_FIXED_ETA": 0.05,
        "PT_MAX_ROUNDS": 200,
        "PT_EARLY_STOP": 30,
        "PT_PREHEAT_ROUNDS": 15,
        "PT_DYNAMIC_INTERVAL": 5,

        "KWH_FIXED_DEPTH": 6,
        "KWH_FIXED_GAMMA": 1.0,
        "KWH_FIXED_ETA": 0.05,
        "KWH_FIXED_ALPHA": 1.0,
        "KWH_ALPHA_MODE": "dynamic",
        "KWH_MAX_ROUNDS": 200,
        "KWH_EARLY_STOP": 30,
        "KWH_PREHEAT_ROUNDS": 15,
        "KWH_DYNAMIC_INTERVAL": 5,

        "DEPTH_DELTA_POS": 1.5,
        "DEPTH_DELTA_NEG": 1.2,
        "DEPTH_SKEW_COEF": 5.0,
        "DEPTH_KURT_COEF": 5.0,
        "DEPTH_TRIM_RATIO": 0.05,
        "DEPTH_SMOOTH_FACTOR": 0.3,
        "DEPTH_MIN": 2,
        "DEPTH_MAX": 20,

        "GAMMA_BASE": 1.0,
        "GAMMA_ALPHA": 0.7,
        "GAMMA_SKEW_COEF": 5.0,
        "GAMMA_KURT_COEF": 5.0,
        "GAMMA_TRIM_RATIO": 0.05,
        "GAMMA_SMOOTH_FACTOR": 0.3,
        "GAMMA_MIN": 0.001,
        "GAMMA_MAX": 20.0,

        "HUBER_SCALE": 2.0,
    }


def get_run_plan(run_mode):
    if run_mode == "screening":
        return {
            "windows": [60, 120, 240],
            "n_runs": 1
        }
    elif run_mode == "confirm":
        return {
            "windows": [30, 60, 120, 240, 360, 480],
            "n_runs": 1
        }
    else:
        raise ValueError("RUN_MODE must be either “screening” or “confirm”")


def get_sensitivity_plan():
    return {
        "SKEW_COEF": [1, 3, 5, 7, 9],
        "KURT_COEF": [1, 3, 5, 7, 9],
        "PREHEAT_ROUNDS": [5, 10, 15, 20, 25],
        "GAMMA_ALPHA": [0.1, 0.3, 0.5, 0.7, 0.9],
    }


def get_baseline_value(param_name, base_cfg):
    if param_name == "SKEW_COEF":
        return base_cfg["DEPTH_SKEW_COEF"]
    if param_name == "KURT_COEF":
        return base_cfg["DEPTH_KURT_COEF"]
    if param_name == "PREHEAT_ROUNDS":
        return base_cfg["PT_PREHEAT_ROUNDS"]
    if param_name == "GAMMA_ALPHA":
        return base_cfg["GAMMA_ALPHA"]
    return base_cfg.get(param_name, None)


def apply_sensitivity_value(cfg, param_name, value):
    if param_name == "SKEW_COEF":
        cfg["DEPTH_SKEW_COEF"] = float(value)
        cfg["GAMMA_SKEW_COEF"] = float(value)
    elif param_name == "KURT_COEF":
        cfg["DEPTH_KURT_COEF"] = float(value)
        cfg["GAMMA_KURT_COEF"] = float(value)
    elif param_name == "PREHEAT_ROUNDS":
        cfg["PT_PREHEAT_ROUNDS"] = int(value)
        cfg["KWH_PREHEAT_ROUNDS"] = int(value)
    elif param_name == "GAMMA_ALPHA":
        cfg["GAMMA_ALPHA"] = float(value)
    else:
        cfg[param_name] = value

    cfg["SENSITIVITY_PARAM_NAME"] = param_name
    cfg["SENSITIVITY_PARAM_VALUE"] = value
    return cfg


def is_baseline_value(param_name, value, base_cfg):
    base_value = get_baseline_value(param_name, base_cfg)
    try:
        return np.isclose(float(value), float(base_value), rtol=0.0, atol=1e-12)
    except Exception:
        return value == base_value


def build_experiments(base_cfg, sensitivity_plan, enable_sensitivity=True):
    experiments = []

    if not enable_sensitivity:
        return experiments

    for param_name, candidate_list in sensitivity_plan.items():
        for value in candidate_list:
            if is_baseline_value(param_name, value, base_cfg):
                continue

            cfg = copy.deepcopy(base_cfg)
            cfg = apply_sensitivity_value(cfg, param_name, value)
            experiments.append({
                "exp_name": f"{param_name}={value}",
                "param_name": param_name,
                "param_value": value,
                "cfg": cfg
            })

    return experiments


def get_given_baseline_metrics():
    windows = [30, 60, 120, 240, 360, 480]
    pt_smape = [11.02, 10.48, 10.41, 10.29, 9.68, 10.57]
    pt_mae = [0.95, 0.88, 0.87, 0.87, 0.82, 0.92]
    kwh_smape = [9.06, 7.43, 6.28, 5.75, 6.14, 5.57]
    kwh_mae = [1.79, 1.50, 1.25, 1.16, 1.25, 1.13]

    rows = []
    for w, mpt, spt, mk, sk in zip(windows, pt_mae, pt_smape, kwh_mae, kwh_smape):
        rows.append({
            "window": w,
            "mae_pt": mpt,
            "smape_pt": spt,
            "mae_kWh": mk,
            "smape_kWh": sk,
        })
    return rows


def fmt2(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return ""
    return f"{float(x):.2f}"


def fmt_mean_pm_maxdiff(values):
    arr = np.array([float(v) for v in values if pd.notna(v)], dtype=float)
    if len(arr) == 0:
        return ""
    mean_val = np.mean(arr)
    plus_val = np.max(arr) - mean_val
    return f"{mean_val:.2f}±{plus_val:.2f}"


def select_param_value_rows(df, param_name, value):
    if df is None or df.empty:
        return pd.DataFrame()

    mask_param = df["param_name"].astype(str) == str(param_name)

    value_series_num = pd.to_numeric(df["param_value"], errors="coerce")
    try:
        value_num = float(value)
        mask_value_num = value_series_num.notna() & np.isclose(value_series_num.astype(float), value_num, rtol=0.0, atol=1e-12)
    except Exception:
        mask_value_num = pd.Series(False, index=df.index)

    mask_value_str = df["param_value"].astype(str) == str(value)
    return df[mask_param & (mask_value_num | mask_value_str)].copy()


def validate_formatted_output(formatted_df, sensitivity_plan):
    problems = []
    for param_name, candidate_list in sensitivity_plan.items():
        for value in candidate_list:
            sub = select_param_value_rows(formatted_df, param_name, value)
            n_window = int((sub["row_type"].astype(str) == "window").sum()) if not sub.empty else 0
            n_avg = int((sub["row_type"].astype(str) == "average").sum()) if not sub.empty else 0
            if n_window != 6 or n_avg != 1:
                problems.append(f"{param_name}={value}: window_rows={n_window}, avg_rows={n_avg}")

    if problems:
        print("\n[WARNING] sensitivity_raw_metrics.csv is parameter groups that have not been fully written:")
        for item in problems:
            print("  - " + item)
        print("Please check whether the corresponding experiment was interrupted or whether this window was skipped because the sample was empty.")
    else:
        print("\n[OK] All candidate values for the four sensitivity parameters have been written to 6 window rows + 1 AVG row.")


def save_formatted_metrics_xlsx(formatted_df, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except Exception as e:
        print(f"openpyxl is not installed; skipping xlsx formatting output: {e}")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "sensitivity_raw_metrics"

    columns = list(formatted_df.columns)
    ws.append(columns)

    metric_cols = {"mae_pt", "smape_pt", "mae_kWh", "smape_kWh"}
    metric_col_idx = {columns.index(c) + 1 for c in metric_cols if c in columns}

    for _, r in formatted_df.iterrows():
        row_values = []
        is_avg = str(r.get("row_type", "")) == "average"
        for c in columns:
            v = r[c]
            if c in metric_cols and not is_avg and v != "":
                try:
                    row_values.append(float(v))
                except Exception:
                    row_values.append(v)
            else:
                row_values.append(v)
        ws.append(row_values)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        row_type = row[columns.index("row_type")].value if "row_type" in columns else ""
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
        if row_type == "average":
            for cell in row:
                cell.font = Font(bold=True)
        else:
            for idx in metric_col_idx:
                row[idx - 1].number_format = "0.00"

    for col_i, col_name in enumerate(columns, start=1):
        width = max(12, min(28, max(len(str(col_name)), 12)))
        ws.column_dimensions[get_column_letter(col_i)].width = width

    wb.save(output_path)


def build_numeric_rows_with_baseline(metrics_df, sensitivity_plan, base_cfg):
    all_rows = []
    baseline_rows = get_given_baseline_metrics()

    if metrics_df is None or metrics_df.empty:
        metrics_df = pd.DataFrame()

    for param_name, candidate_list in sensitivity_plan.items():
        for value in candidate_list:
            exp_name = f"{param_name}={value}"

            if is_baseline_value(param_name, value, base_cfg):
                for item in baseline_rows:
                    row = {
                        "exp_name": exp_name,
                        "param_name": param_name,
                        "param_value": value,
                        "run": 1,
                        "window": item["window"],
                        "mae_pt": item["mae_pt"],
                        "smape_pt": item["smape_pt"],
                        "rounds_pt": np.nan,
                        "best_round_pt": np.nan,
                        "time_pt_ms": np.nan,
                        "mae_kWh": item["mae_kWh"],
                        "smape_kWh": item["smape_kWh"],
                        "rounds_kWh": np.nan,
                        "best_round_kWh": np.nan,
                        "time_kWh_ms": np.nan,
                        "window_time_ms": np.nan,
                        "is_baseline_direct": True,
                    }
                    all_rows.append(row)
            else:
                if metrics_df.empty:
                    continue
                sub = select_param_value_rows(metrics_df, param_name, value)
                if sub.empty:
                    print(f"[WARNING] No results found for: {param_name}={value}; this parameter value will not be written to the final table.")
                    continue
                sub = sub.sort_values(["run", "window"])
                for _, r in sub.iterrows():
                    row = r.to_dict()
                    row["is_baseline_direct"] = False
                    all_rows.append(row)

    return pd.DataFrame(all_rows)


def build_formatted_raw_metrics(numeric_df, sensitivity_plan):
    formatted_rows = []
    metric_cols = ["mae_pt", "smape_pt", "mae_kWh", "smape_kWh"]

    for param_name, candidate_list in sensitivity_plan.items():
        for value in candidate_list:
            exp_name = f"{param_name}={value}"
            sub = select_param_value_rows(numeric_df, param_name, value)
            if sub.empty:
                continue

            sub = sub[pd.to_numeric(sub["window"], errors="coerce").notna()].copy()
            sub["window"] = pd.to_numeric(sub["window"], errors="coerce").astype(int)
            sub = sub.sort_values("window")

            for _, r in sub.iterrows():
                formatted_rows.append({
                    "exp_name": exp_name,
                    "param_name": param_name,
                    "param_value": value,
                    "run": int(r["run"]) if pd.notna(r.get("run", np.nan)) else "",
                    "window": int(r["window"]) if pd.notna(r.get("window", np.nan)) else "",
                    "mae_pt": fmt2(r["mae_pt"]),
                    "smape_pt": fmt2(r["smape_pt"]),
                    "rounds_pt": int(r["rounds_pt"]) if pd.notna(r.get("rounds_pt", np.nan)) else "",
                    "best_round_pt": int(r["best_round_pt"]) if pd.notna(r.get("best_round_pt", np.nan)) else "",
                    "time_pt_ms": fmt2(r["time_pt_ms"]) if pd.notna(r.get("time_pt_ms", np.nan)) else "",
                    "mae_kWh": fmt2(r["mae_kWh"]),
                    "smape_kWh": fmt2(r["smape_kWh"]),
                    "rounds_kWh": int(r["rounds_kWh"]) if pd.notna(r.get("rounds_kWh", np.nan)) else "",
                    "best_round_kWh": int(r["best_round_kWh"]) if pd.notna(r.get("best_round_kWh", np.nan)) else "",
                    "time_kWh_ms": fmt2(r["time_kWh_ms"]) if pd.notna(r.get("time_kWh_ms", np.nan)) else "",
                    "window_time_ms": fmt2(r["window_time_ms"]) if pd.notna(r.get("window_time_ms", np.nan)) else "",
                    "row_type": "window",
                    "is_baseline_direct": bool(r.get("is_baseline_direct", False)),
                })

            avg_row = {
                "exp_name": exp_name,
                "param_name": param_name,
                "param_value": value,
                "run": "",
                "window": "AVG",
                "rounds_pt": "",
                "best_round_pt": "",
                "time_pt_ms": "",
                "rounds_kWh": "",
                "best_round_kWh": "",
                "time_kWh_ms": "",
                "window_time_ms": "",
                "row_type": "average",
                "is_baseline_direct": bool(sub["is_baseline_direct"].all()) if "is_baseline_direct" in sub else False,
            }
            for c in metric_cols:
                avg_row[c] = fmt_mean_pm_maxdiff(sub[c].values)

            formatted_rows.append(avg_row)

    columns = [
        "exp_name", "param_name", "param_value", "run", "window",
        "mae_pt", "smape_pt", "rounds_pt", "best_round_pt", "time_pt_ms",
        "mae_kWh", "smape_kWh", "rounds_kWh", "best_round_kWh", "time_kWh_ms",
        "window_time_ms", "row_type", "is_baseline_direct"
    ]
    return pd.DataFrame(formatted_rows, columns=columns)


def build_preprocessed_dataframe(csv_path):
    df0 = pd.read_csv(csv_path, parse_dates=["connection_time_copy"])

    df0 = df0[
        (df0.parking_time <= df0.Requested_parking_time + 2) &
        (df0.kWhRequested <= 150) &
        (df0.kWhDelivered <= df0.kWhRequested)
    ].dropna(subset=["connection_time_copy"]).reset_index(drop=True)

    df0["hour"] = df0.connection_time_copy.dt.hour
    df0["weekday"] = df0.connection_time_copy.dt.weekday
    df0["month"] = df0.connection_time_copy.dt.month
    df0["dayofyear"] = df0.connection_time_copy.dt.dayofyear

    for c in ["hour", "weekday", "month", "dayofyear"]:
        denom = max(df0[c].max(), 1)
        df0[f"{c}_sin"] = np.sin(2 * np.pi * df0[c] / denom)
        df0[f"{c}_cos"] = np.cos(2 * np.pi * df0[c] / denom)

    for c in ["userID", "stationID"]:
        le = LabelEncoder()
        df0[c] = le.fit_transform(df0[c].astype(str))

    df0.sort_values("connection_time_copy", inplace=True)
    df0.reset_index(drop=True, inplace=True)
    return df0


def main():
    base_cfg = get_base_config()
    run_plan = get_run_plan(base_cfg["RUN_MODE"])
    sensitivity_plan = get_sensitivity_plan()
    experiments = build_experiments(
        base_cfg,
        sensitivity_plan,
        enable_sensitivity=base_cfg["ENABLE_SENSITIVITY"]
    )

    os.makedirs(base_cfg["RESULTS_ROOT"], exist_ok=True)

    with open(os.path.join(base_cfg["RESULTS_ROOT"], "base_config.json"), "w", encoding="utf-8") as f:
        json.dump(base_cfg, f, indent=2, ensure_ascii=False)

    df0 = build_preprocessed_dataframe(base_cfg["CSV_PATH"])

    test_start = datetime.strptime(base_cfg["TEST_START"], "%Y-%m-%d")
    test_end = datetime.strptime(base_cfg["TEST_END"], "%Y-%m-%d")
    eps = base_cfg["EPS"]

    metrics_all = []
    trace_all = []

    print(f"Total experiments: {len(experiments)}")
    print("Sensitivity plan: SKEW_COEF / KURT_COEF / PREHEAT_ROUNDS / GAMMA_ALPHA")
    print(f"Run mode: {base_cfg['RUN_MODE']}")
    print(f"Windows: {base_cfg['WINDOWS']}")
    print(f"n_runs: {run_plan['n_runs']}")

    for exp in experiments:
        exp_name = exp["exp_name"]
        param_name = exp["param_name"]
        param_value = exp["param_value"]
        cfg = exp["cfg"]

        safe_exp_name = make_safe_name(exp_name)
        exp_dir = os.path.join(base_cfg["RESULTS_ROOT"], safe_exp_name)
        os.makedirs(exp_dir, exist_ok=True)

        with open(os.path.join(exp_dir, "experiment_config.json"), "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2, ensure_ascii=False)

        print("\n" + "=" * 80)
        print(f"Experiment: {exp_name}")
        print("=" * 80)

        for run in range(1, run_plan["n_runs"] + 1):
            run_dir = os.path.join(exp_dir, f"run_{run}")
            curve_dir_pt = os.path.join(run_dir, "parking_time")
            curve_dir_kwh = os.path.join(run_dir, "kWhDelivered")
            os.makedirs(curve_dir_pt, exist_ok=True)
            os.makedirs(curve_dir_kwh, exist_ok=True)

            print(f"\n=== Experiment {exp_name} | Run {run} ===")

            df = df0.copy()

            hist_all = df[df.connection_time_copy < test_start]
            if hist_all.empty:
                print("No history data before test_start, skip this run.")
                continue

            ua_num = hist_all.groupby("userID").kWhDelivered.mean()
            ua_den = hist_all.groupby("userID").kWhRequested.mean() + 1e-6
            ua = ua_num / ua_den

            if len(ua) == 0:
                df["cluster"] = 0
            else:
                n_clusters = min(10, len(ua))
                km = KMeans(n_clusters=n_clusters, random_state=cfg["BASE_SEED"] + run)
                km.fit(ua.values.reshape(-1, 1))
                cluster_map = dict(zip(ua.index, km.labels_))
                df["cluster"] = df.userID.map(cluster_map).fillna(0).astype(int)

            run_summary = []

            for w in base_cfg["WINDOWS"]:
                print(f"\n>>> Window = {w} days")
                t_window_start = time.perf_counter()

                tr0 = test_start - timedelta(days=w)
                tr = df[(df.connection_time_copy >= tr0) & (df.connection_time_copy < test_start)].copy()
                te = df[(df.connection_time_copy >= test_start) & (df.connection_time_copy <= test_end)].copy()

                if tr.empty or te.empty:
                    print("Training or testing set empty, skip.")
                    continue

                threshold_map = {}
                for col in ["parking_time", "kWhDelivered"]:
                    tr[col], threshold_map[col] = wavelet_denoise_series(
                        tr[col].values, wavelet="db4", level=3, thr=None
                    )

                if len(tr) >= 5:
                    lof_neighbors = min(20, len(tr) - 1)
                    if lof_neighbors >= 2:
                        mask = LocalOutlierFactor(n_neighbors=lof_neighbors).fit_predict(
                            tr[["parking_time", "kWhDelivered"]]
                        )
                        tr = tr[mask == 1].reset_index(drop=True)

                for col in ["parking_time", "kWhDelivered"]:
                    te[col], _ = wavelet_denoise_series(
                        te[col].values, wavelet="db4", level=3, thr=threshold_map[col]
                    )

                tr["age_days"] = (test_start - tr.connection_time_copy).dt.days
                w_decay = np.exp(-0.015 * tr["age_days"])

                tr.reset_index(drop=True, inplace=True)
                te.reset_index(drop=True, inplace=True)

                tr["is_weekend"] = tr.weekday.isin([5, 6]).astype(int)
                te["is_weekend"] = te.weekday.isin([5, 6]).astype(int)
                tr["is_holiday"] = tr.get("connectionTime_is_holiday", 0)
                te["is_holiday"] = te.get("connectionTime_is_holiday", 0)

                sim_feats = ["hour", "weekday", "kWhRequested", "Requested_parking_time"]
                ss = StandardScaler().fit(tr[sim_feats])
                tr_s = ss.transform(tr[sim_feats])
                te_s = ss.transform(te[sim_feats])

                tr["sim_park"] = 0.0
                tr["sim_kwh"] = 0.0

                for i in range(len(tr)):
                    prev_idx = tr[(tr.userID == tr.loc[i, "userID"]) & (tr.index < i)].index
                    if len(prev_idx) == 0:
                        tr.at[i, "sim_park"] = tr.parking_time.median()
                        tr.at[i, "sim_kwh"] = tr.kWhDelivered.median()
                    else:
                        sims = np.dot(tr_s[prev_idx], tr_s[i])
                        sort_idx = np.argsort(sims)[-5:]
                        top_idx = prev_idx[sort_idx]
                        w5 = sims[sort_idx]
                        tr.at[i, "sim_park"] = np.average(
                            tr.loc[top_idx, "parking_time"], weights=w5 + 1e-6
                        )
                        tr.at[i, "sim_kwh"] = np.average(
                            tr.loc[top_idx, "kWhDelivered"], weights=w5 + 1e-6
                        )

                te["sim_park"] = 0.0
                te["sim_kwh"] = 0.0

                for i in range(len(te)):
                    prev_idx = tr[tr.userID == te.loc[i, "userID"]].index
                    if len(prev_idx) == 0:
                        te.at[i, "sim_park"] = tr.parking_time.median()
                        te.at[i, "sim_kwh"] = tr.kWhDelivered.median()
                    else:
                        sims = np.dot(tr_s[prev_idx], te_s[i])
                        sort_idx = np.argsort(sims)[-5:]
                        top_idx = prev_idx[sort_idx]
                        w5 = sims[sort_idx]
                        te.at[i, "sim_park"] = np.average(
                            tr.loc[top_idx, "parking_time"], weights=w5 + 1e-6
                        )
                        te.at[i, "sim_kwh"] = np.average(
                            tr.loc[top_idx, "kWhDelivered"], weights=w5 + 1e-6
                        )

                ua_m = tr.groupby("userID").kWhDelivered.mean()
                ua_c = tr.userID.value_counts()
                sa_m = tr.groupby("stationID").kWhDelivered.mean()
                sa_c = tr.stationID.value_counts()

                for D in (tr, te):
                    D["user_avg_kWh"] = D.userID.map(ua_m).fillna(ua_m.mean())
                    D["user_freq"] = D.userID.map(ua_c).fillna(0)
                    D["station_avg_kWh"] = D.stationID.map(sa_m).fillna(sa_m.mean())
                    D["station_freq"] = D.stationID.map(sa_c).fillna(0)

                tr["hour_of_week"] = tr.weekday * 24 + tr.hour
                te["hour_of_week"] = te.weekday * 24 + te.hour

                tr["req_rate"] = tr.kWhRequested / (tr.Requested_parking_time + 1e-6)
                te["req_rate"] = te.kWhRequested / (te.Requested_parking_time + 1e-6)

                sh = tr.groupby(["stationID", "hour_of_week"]).kWhDelivered.sum() / \
                     (tr.groupby(["stationID", "hour_of_week"]).kWhRequested.sum() + 1e-6)
                uh = tr.groupby(["userID", "hour_of_week"]).kWhDelivered.sum() / \
                     (tr.groupby(["userID", "hour_of_week"]).kWhRequested.sum() + 1e-6)
                shu = tr.groupby(["stationID", "userID", "hour_of_week"]).kWhDelivered.sum() / \
                      (tr.groupby(["stationID", "userID", "hour_of_week"]).kWhRequested.sum() + 1e-6)

                tr["sh_ratio"] = tr.set_index(["stationID", "hour_of_week"]).index.map(sh)
                tr["uh_ratio"] = tr.set_index(["userID", "hour_of_week"]).index.map(uh)
                tr["shur_ratio"] = tr.set_index(["stationID", "userID", "hour_of_week"]).index.map(shu)

                tr["sh_ratio"] = tr["sh_ratio"].fillna(sh.mean())
                tr["uh_ratio"] = tr["uh_ratio"].fillna(uh.mean())
                tr["shur_ratio"] = tr["shur_ratio"].fillna(shu.mean())

                te["sh_ratio"] = te.set_index(["stationID", "hour_of_week"]).index.map(sh).fillna(sh.mean())
                te["uh_ratio"] = te.set_index(["userID", "hour_of_week"]).index.map(uh).fillna(uh.mean())
                te["shur_ratio"] = te.set_index(["stationID", "userID", "hour_of_week"]).index.map(shu).fillna(shu.mean())

                pf = [
                    "hour", "weekday", "month", "is_weekend", "is_holiday",
                    "hour_sin", "hour_cos", "weekday_sin", "weekday_cos",
                    "month_sin", "month_cos", "userID", "stationID", "sim_park"
                ]
                Xp_tr = tr[pf].values
                y_p_tr = tr.parking_time.values
                Xp_te = te[pf].values
                y_p_te = te.parking_time.values

                n_val = max(int(0.2 * len(tr)), 50)
                n_val = min(n_val, len(tr) - 1)
                if len(tr) - n_val < 10:
                    n_val = max(1, len(tr) // 5)
                if len(tr) - n_val < 5:
                    print("Too few training samples after split, skip window.")
                    continue

                Xtr_p, Xva_p = Xp_tr[:-n_val], Xp_tr[-n_val:]
                ytr_p, yva_p = y_p_tr[:-n_val], y_p_tr[-n_val:]
                wtr_p = w_decay[:-n_val]

                dtrain_p = xgb.DMatrix(Xtr_p, label=ytr_p, weight=wtr_p)
                dval_p = xgb.DMatrix(Xva_p, label=yva_p)

                pt_val_maes = []
                model_p = None
                best_model_p = None
                best_va_p = np.inf
                best_round_p = 0
                rounds_no_improve_p = 0

                prev_depth_p = None
                prev_gamma_p = None
                depth_p = cfg["PT_FIXED_DEPTH"]
                gamma_p = cfg["PT_FIXED_GAMMA"]

                start_p = time.perf_counter()
                for m in range(cfg["PT_MAX_ROUNDS"]):
                    updated_flag = 0
                    depth_info_p = {
                        "sigma_w": np.nan, "skew_val": np.nan, "kurt_val": np.nan,
                        "mad_w": np.nan, "f": np.nan, "raw": np.nan
                    }
                    gamma_info_p = {
                        "iqr_w": np.nan, "mad_w": np.nan, "skew_val": np.nan,
                        "kurt_val": np.nan, "f_gamma": np.nan, "raw": np.nan
                    }

                    if m < cfg["PT_PREHEAT_ROUNDS"]:
                        depth_p = cfg["PT_FIXED_DEPTH"]
                        gamma_p = cfg["PT_FIXED_GAMMA"]
                        stage_p = "preheat"
                    else:
                        stage_p = "dynamic"
                        if (m - cfg["PT_PREHEAT_ROUNDS"]) % cfg["PT_DYNAMIC_INTERVAL"] == 0:
                            updated_flag = 1
                            if model_p is None:
                                residuals_p = ytr_p.copy()
                            else:
                                residuals_p = ytr_p - model_p.predict(dtrain_p)

                            depth_p, depth_info_p = compute_dynamic_max_depth_advanced(
                                residuals_p, wtr_p,
                                d0=cfg["PT_FIXED_DEPTH"],
                                delta_pos=cfg["DEPTH_DELTA_POS"],
                                delta_neg=cfg["DEPTH_DELTA_NEG"],
                                skew_coef=cfg["DEPTH_SKEW_COEF"],
                                kurt_coef=cfg["DEPTH_KURT_COEF"],
                                min_depth=cfg["DEPTH_MIN"],
                                max_depth_limit=cfg["DEPTH_MAX"],
                                prev_depth=prev_depth_p,
                                trim_ratio=cfg["DEPTH_TRIM_RATIO"],
                                smooth_factor=cfg["DEPTH_SMOOTH_FACTOR"],
                                return_detail=True
                            )

                            gamma_p, gamma_info_p = compute_dynamic_gamma_advanced(
                                residuals_p, wtr_p,
                                base=cfg["GAMMA_BASE"],
                                gamma_alpha=cfg["GAMMA_ALPHA"],
                                skew_coef=cfg["GAMMA_SKEW_COEF"],
                                kurt_coef=cfg["GAMMA_KURT_COEF"],
                                min_gamma=cfg["GAMMA_MIN"],
                                max_gamma=cfg["GAMMA_MAX"],
                                prev_gamma=prev_gamma_p,
                                trim_ratio=cfg["GAMMA_TRIM_RATIO"],
                                smooth_factor=cfg["GAMMA_SMOOTH_FACTOR"],
                                return_detail=True
                            )

                            prev_depth_p = depth_p
                            prev_gamma_p = gamma_p

                    params_p = {
                        "tree_method": "hist",
                        "eta": cfg["PT_FIXED_ETA"],
                        "max_depth": depth_p,
                        "gamma": gamma_p,
                        "objective": "reg:squarederror",
                        "eval_metric": "mae",
                        "verbosity": 0,
                        "subsample": 0.8,
                        "colsample_bytree": 0.8,
                        "seed": cfg["BASE_SEED"] + run * 1000 + m
                    }

                    if model_p is None:
                        model_p = xgb.train(params_p, dtrain_p, num_boost_round=1, verbose_eval=False)
                    else:
                        model_p = xgb.train(
                            params_p, dtrain_p, num_boost_round=1,
                            xgb_model=model_p, verbose_eval=False
                        )

                    pred_va_p = model_p.predict(dval_p)
                    va_mae_p = mae(yva_p, pred_va_p)
                    pt_val_maes.append(va_mae_p)

                    if va_mae_p < best_va_p - 1e-6:
                        best_va_p = va_mae_p
                        best_round_p = m + 1
                        rounds_no_improve_p = 0
                        best_model_p = model_p.copy()
                    else:
                        rounds_no_improve_p += 1

                    trace_all.append({
                        "exp_name": exp_name,
                        "param_name": param_name,
                        "param_value": param_value,
                        "run": run,
                        "window": w,
                        "task": "parking_time",
                        "round": m + 1,
                        "stage": stage_p,
                        "updated_flag": updated_flag,
                        "depth": depth_p,
                        "gamma": gamma_p,
                        "huber_alpha": np.nan,
                        "depth_sigma_w": depth_info_p["sigma_w"],
                        "depth_skew": depth_info_p["skew_val"],
                        "depth_kurt": depth_info_p["kurt_val"],
                        "depth_mad_w": depth_info_p["mad_w"],
                        "depth_f": depth_info_p["f"],
                        "depth_raw": depth_info_p["raw"],
                        "gamma_iqr_w": gamma_info_p["iqr_w"],
                        "gamma_mad_w": gamma_info_p["mad_w"],
                        "gamma_skew": gamma_info_p["skew_val"],
                        "gamma_kurt": gamma_info_p["kurt_val"],
                        "gamma_f": gamma_info_p["f_gamma"],
                        "gamma_raw": gamma_info_p["raw"],
                        "val_mae": va_mae_p
                    })

                    if rounds_no_improve_p >= cfg["PT_EARLY_STOP"]:
                        break

                end_p = time.perf_counter()
                rounds_p = len(pt_val_maes)
                time_p_ms = (end_p - start_p) * 1000

                final_model_p = best_model_p if best_model_p is not None else model_p
                pred_p_te = final_model_p.predict(xgb.DMatrix(Xp_te))
                mae_p = mae(y_p_te, pred_p_te)
                sm_p = smape(y_p_te, pred_p_te)

                save_convergence_curve(
                    pt_val_maes,
                    title=f"Window={w} parking_time Val MAE vs Rounds",
                    save_path=os.path.join(curve_dir_pt, f"convergence_parking_time_window_{w}.png"),
                    best_round=best_round_p
                )

                pt_transform = PowerTransformer(method='box-cox', standardize=False)
                y_all_k = tr.kWhDelivered.values + eps
                y_bc = pt_transform.fit_transform(y_all_k.reshape(-1, 1)).flatten()

                tr = tr.sort_values("connection_time_copy").reset_index(drop=True)
                te = te.sort_values("connection_time_copy").reset_index(drop=True)

                user_roll3 = tr.groupby("userID")["kWhDelivered"] \
                               .apply(lambda x: x.shift(1).rolling(3, min_periods=1).mean()) \
                               .reset_index(level=0, drop=True)
                tr["user_roll_kWh3"] = user_roll3.fillna(tr.kWhDelivered.median())

                station_roll3 = tr.groupby("stationID")["kWhDelivered"] \
                                  .apply(lambda x: x.shift(1).rolling(3, min_periods=1).mean()) \
                                  .reset_index(level=0, drop=True)
                tr["station_roll_kWh3"] = station_roll3.fillna(tr.kWhDelivered.median())

                user_hist = tr.groupby("userID")["kWhDelivered"].apply(list).to_dict()
                station_hist = tr.groupby("stationID")["kWhDelivered"].apply(list).to_dict()

                te["user_roll_kWh3"] = te["userID"].apply(
                    lambda u: np.mean(user_hist[u][-3:]) if (u in user_hist and len(user_hist[u]) > 0)
                    else tr.kWhDelivered.median()
                )
                te["station_roll_kWh3"] = te["stationID"].apply(
                    lambda s: np.mean(station_hist[s][-3:]) if (s in station_hist and len(station_hist[s]) > 0)
                    else tr.kWhDelivered.median()
                )

                med_k = tr.kWhDelivered.median()
                tr["pk1"] = tr.groupby("userID").kWhDelivered.shift(1).fillna(med_k)
                tr["pk2"] = tr.groupby("userID").kWhDelivered.shift(2).fillna(med_k)
                tr["rk3"] = tr.groupby("userID").kWhDelivered.rolling(3, min_periods=1).mean() \
                             .reset_index(0, drop=True).fillna(med_k)

                te["pk1"] = med_k
                te["pk2"] = med_k
                te["rk3"] = med_k
                prev_vals = []
                for i in range(len(te)):
                    if i >= 1:
                        te.at[i, "pk1"] = prev_vals[i - 1]
                    if i >= 2:
                        te.at[i, "pk2"] = prev_vals[i - 2]
                    if i >= 3:
                        te.at[i, "rk3"] = np.mean(prev_vals[i - 3:i])
                    prev_vals.append(te.at[i, "kWhDelivered"])

                feats_k = [
                    "hour", "weekday", "month", "is_weekend", "is_holiday",
                    "hour_sin", "hour_cos", "weekday_sin", "weekday_cos",
                    "month_sin", "month_cos", "dayofyear_sin", "dayofyear_cos",
                    "cluster", "sim_kwh", "user_avg_kWh", "user_freq",
                    "station_avg_kWh", "station_freq",
                    "req_rate", "hour_of_week",
                    "sh_ratio", "uh_ratio", "shur_ratio",
                    "pk1", "pk2", "rk3",
                    "station_roll_kWh3"
                ]

                X_all_tr_k = tr[feats_k].values
                X_all_te_k = te[feats_k].values

                Xk_tr = X_all_tr_k[:-n_val]
                Xk_va = X_all_tr_k[-n_val:]
                ytr_bc = y_bc[:-n_val]
                yva_bc = y_bc[-n_val:]
                wtr_k = w_decay[:-n_val]

                dtrain_k = xgb.DMatrix(Xk_tr, label=ytr_bc, weight=wtr_k)
                dval_k = xgb.DMatrix(Xk_va, label=yva_bc)

                best_va_k = np.inf
                best_round_k = 0
                kwh_val_maes = []
                model_k = None
                best_model_k = None
                rounds_no_improve_k = 0

                prev_depth_k = None
                prev_gamma_k = None
                depth_k = cfg["KWH_FIXED_DEPTH"]
                gamma_k = cfg["KWH_FIXED_GAMMA"]

                start_k = time.perf_counter()
                for m in range(cfg["KWH_MAX_ROUNDS"]):
                    updated_flag = 0
                    depth_info_k = {
                        "sigma_w": np.nan, "skew_val": np.nan, "kurt_val": np.nan,
                        "mad_w": np.nan, "f": np.nan, "raw": np.nan
                    }
                    gamma_info_k = {
                        "iqr_w": np.nan, "mad_w": np.nan, "skew_val": np.nan,
                        "kurt_val": np.nan, "f_gamma": np.nan, "raw": np.nan
                    }

                    if m < cfg["KWH_PREHEAT_ROUNDS"]:
                        depth_k = cfg["KWH_FIXED_DEPTH"]
                        gamma_k = cfg["KWH_FIXED_GAMMA"]
                        stage_k = "preheat"
                    else:
                        stage_k = "dynamic"
                        if (m - cfg["KWH_PREHEAT_ROUNDS"]) % cfg["KWH_DYNAMIC_INTERVAL"] == 0:
                            updated_flag = 1
                            if model_k is None:
                                residuals_k = ytr_bc.copy()
                            else:
                                residuals_k = ytr_bc - model_k.predict(
                                    xgb.DMatrix(Xk_tr, label=ytr_bc, weight=wtr_k)
                                )

                            depth_k, depth_info_k = compute_dynamic_max_depth_advanced(
                                residuals_k, wtr_k,
                                d0=cfg["KWH_FIXED_DEPTH"],
                                delta_pos=cfg["DEPTH_DELTA_POS"],
                                delta_neg=cfg["DEPTH_DELTA_NEG"],
                                skew_coef=cfg["DEPTH_SKEW_COEF"],
                                kurt_coef=cfg["DEPTH_KURT_COEF"],
                                min_depth=cfg["DEPTH_MIN"],
                                max_depth_limit=cfg["DEPTH_MAX"],
                                prev_depth=prev_depth_k,
                                trim_ratio=cfg["DEPTH_TRIM_RATIO"],
                                smooth_factor=cfg["DEPTH_SMOOTH_FACTOR"],
                                return_detail=True
                            )

                            gamma_k, gamma_info_k = compute_dynamic_gamma_advanced(
                                residuals_k, wtr_k,
                                base=cfg["GAMMA_BASE"],
                                gamma_alpha=cfg["GAMMA_ALPHA"],
                                skew_coef=cfg["GAMMA_SKEW_COEF"],
                                kurt_coef=cfg["GAMMA_KURT_COEF"],
                                min_gamma=cfg["GAMMA_MIN"],
                                max_gamma=cfg["GAMMA_MAX"],
                                prev_gamma=prev_gamma_k,
                                trim_ratio=cfg["GAMMA_TRIM_RATIO"],
                                smooth_factor=cfg["GAMMA_SMOOTH_FACTOR"],
                                return_detail=True
                            )

                            prev_depth_k = depth_k
                            prev_gamma_k = gamma_k

                    if model_k is None:
                        residuals_k = ytr_bc.copy()
                    else:
                        residuals_k = ytr_bc - model_k.predict(
                            xgb.DMatrix(Xk_tr, label=ytr_bc, weight=wtr_k)
                        )

                    if cfg["KWH_ALPHA_MODE"] == "fixed":
                        huber_alpha = cfg["KWH_FIXED_ALPHA"]
                    else:
                        huber_alpha = compute_huber_alpha(residuals_k, scale=cfg["HUBER_SCALE"])

                    params_k = {
                        "tree_method": "hist",
                        "eta": cfg["KWH_FIXED_ETA"],
                        "max_depth": depth_k,
                        "gamma": gamma_k,
                        "objective": "reg:pseudohubererror",
                        "huber_slope": huber_alpha,
                        "eval_metric": "mae",
                        "verbosity": 0,
                        "seed": cfg["BASE_SEED"] + run * 2000 + m
                    }

                    if model_k is None:
                        model_k = xgb.train(params_k, dtrain_k, num_boost_round=1, verbose_eval=False)
                    else:
                        model_k = xgb.train(
                            params_k, dtrain_k, num_boost_round=1,
                            xgb_model=model_k, verbose_eval=False
                        )

                    pred_va_k_now = model_k.predict(dval_k)
                    va_k_mae = mae(yva_bc, pred_va_k_now)
                    kwh_val_maes.append(va_k_mae)

                    if va_k_mae < best_va_k - 1e-6:
                        best_va_k = va_k_mae
                        best_round_k = m + 1
                        rounds_no_improve_k = 0
                        best_model_k = model_k.copy()
                    else:
                        rounds_no_improve_k += 1

                    trace_all.append({
                        "exp_name": exp_name,
                        "param_name": param_name,
                        "param_value": param_value,
                        "run": run,
                        "window": w,
                        "task": "kWhDelivered",
                        "round": m + 1,
                        "stage": stage_k,
                        "updated_flag": updated_flag,
                        "depth": depth_k,
                        "gamma": gamma_k,
                        "huber_alpha": huber_alpha,
                        "depth_sigma_w": depth_info_k["sigma_w"],
                        "depth_skew": depth_info_k["skew_val"],
                        "depth_kurt": depth_info_k["kurt_val"],
                        "depth_mad_w": depth_info_k["mad_w"],
                        "depth_f": depth_info_k["f"],
                        "depth_raw": depth_info_k["raw"],
                        "gamma_iqr_w": gamma_info_k["iqr_w"],
                        "gamma_mad_w": gamma_info_k["mad_w"],
                        "gamma_skew": gamma_info_k["skew_val"],
                        "gamma_kurt": gamma_info_k["kurt_val"],
                        "gamma_f": gamma_info_k["f_gamma"],
                        "gamma_raw": gamma_info_k["raw"],
                        "val_mae": va_k_mae
                    })

                    if rounds_no_improve_k >= cfg["KWH_EARLY_STOP"]:
                        break

                end_k = time.perf_counter()
                rounds_k = len(kwh_val_maes)
                time_k_ms = (end_k - start_k) * 1000

                final_model_k = best_model_k if best_model_k is not None else model_k

                probe_model_k = model_k

                pred_te_k_xgb_probe_bc = probe_model_k.predict(xgb.DMatrix(X_all_te_k))
                pred_te_k_xgb_probe = pt_transform.inverse_transform(
                    np.maximum(pred_te_k_xgb_probe_bc.reshape(-1, 1), 1e-6)
                ).flatten() - eps
                pred_te_k_xgb_probe = np.maximum(pred_te_k_xgb_probe, 0.0)

                mae_k_xgb_probe = mae(te.kWhDelivered.values, pred_te_k_xgb_probe)
                sm_k_xgb_probe = smape(te.kWhDelivered.values, pred_te_k_xgb_probe)

                pred_tr_k_full = final_model_k.predict(xgb.DMatrix(Xk_tr, label=ytr_bc, weight=wtr_k))
                res_tr_bc = ytr_bc - pred_tr_k_full

                Xk_tr_df = pd.DataFrame(Xk_tr, columns=feats_k)
                Xk_va_df = pd.DataFrame(Xk_va, columns=feats_k)
                X_all_te_df = pd.DataFrame(X_all_te_k, columns=feats_k)

                lgb_params = get_lgb_params(random_state=cfg["BASE_SEED"] + run)
                model_lgb = lgb.LGBMRegressor(**lgb_params)
                model_lgb.fit(Xk_tr_df, res_tr_bc, sample_weight=wtr_k)

                pred_te_k_huber = final_model_k.predict(xgb.DMatrix(X_all_te_k))
                res_te_bc = model_lgb.predict(X_all_te_df)
                final_huber = pred_te_k_huber + res_te_bc

                clipped = np.maximum(final_huber.reshape(-1, 1), 1e-6)
                inv_te_k = pt_transform.inverse_transform(clipped).flatten() - eps

                order = np.argsort(te.connection_time_copy.values)
                smooth = medfilt(inv_te_k[order], kernel_size=5)
                final_k = np.empty_like(smooth)
                final_k[order] = smooth

                mae_k = mae(te.kWhDelivered.values, final_k)
                sm_k = smape(te.kWhDelivered.values, final_k)

                save_convergence_curve(
                    kwh_val_maes,
                    title=f"Window={w} kWhDelivered Val MAE vs Rounds",
                    save_path=os.path.join(curve_dir_kwh, f"convergence_kWh_window_{w}.png"),
                    best_round=best_round_k
                )

                t_window_end = time.perf_counter()
                t_window_ms = (t_window_end - t_window_start) * 1000

                print(f"  Window={w}:")
                print(f"    parking_time -> MAE={mae_p:.3f}, SMAPE={sm_p:.3f}%, rounds={rounds_p}, best_round={best_round_p}, time={time_p_ms:.0f}ms")
                print(f"    kWhDelivered -> MAE={mae_k:.3f}, SMAPE={sm_k:.3f}%, rounds={rounds_k}, best_round={best_round_k}, time={time_k_ms:.0f}ms")

                run_summary.append((
                    w,
                    mae_p, sm_p, rounds_p, best_round_p, time_p_ms,
                    mae_k, sm_k, rounds_k, best_round_k, time_k_ms,
                    t_window_ms
                ))

                metrics_all.append({
                    "exp_name": exp_name,
                    "param_name": param_name,
                    "param_value": param_value,
                    "run": run,
                    "window": w,
                    "mae_pt": mae_p,
                    "smape_pt": sm_p,
                    "rounds_pt": rounds_p,
                    "best_round_pt": best_round_p,
                    "time_pt_ms": time_p_ms,

                    "mae_kWh": mae_k,
                    "smape_kWh": sm_k,

                    "mae_kWh_xgb_probe": mae_k_xgb_probe,
                    "smape_kWh_xgb_probe": sm_k_xgb_probe,

                    "rounds_kWh": rounds_k,
                    "best_round_kWh": best_round_k,
                    "time_kWh_ms": time_k_ms,
                    "window_time_ms": t_window_ms
                })

            if run_summary:
                print(f"\nResults for {exp_name} | run {run}:")
                print("window | MAE_pt | SMAPE_pt | rounds_pt | best_pt | MAE_kWh | SMAPE_kWh | rounds_kWh | best_kWh")
                for item in run_summary:
                    print(
                        f"{item[0]:>6} | "
                        f"{item[1]:6.3f} | {item[2]:7.3f}% | {item[3]:9d} | {item[4]:7d} | "
                        f"{item[6]:7.3f} | {item[7]:9.3f}% | {item[8]:10d} | {item[9]:8d}"
                    )

    metrics_df = pd.DataFrame(metrics_all)
    trace_df = pd.DataFrame(trace_all)

    numeric_run_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_raw_metrics_numeric_without_baseline.csv")
    metrics_df.to_csv(numeric_run_path, index=False)

    numeric_with_baseline_df = build_numeric_rows_with_baseline(metrics_df, sensitivity_plan, base_cfg)
    numeric_with_baseline_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_raw_metrics_numeric.csv")
    numeric_with_baseline_df.to_csv(numeric_with_baseline_path, index=False)

    formatted_raw_df = build_formatted_raw_metrics(numeric_with_baseline_df, sensitivity_plan)
    validate_formatted_output(formatted_raw_df, sensitivity_plan)

    metrics_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_raw_metrics.csv")
    formatted_raw_df.to_csv(metrics_path, index=False, encoding="utf-8-sig")

    metrics_xlsx_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_raw_metrics.xlsx")
    save_formatted_metrics_xlsx(formatted_raw_df, metrics_xlsx_path)

    trace_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_dynamic_trace.csv")
    trace_df.to_csv(trace_path, index=False)

    summary_df = numeric_with_baseline_df.groupby(["param_name", "param_value"], dropna=False).agg(
        mae_pt_mean=("mae_pt", "mean"),
        mae_pt_max=("mae_pt", "max"),
        smape_pt_mean=("smape_pt", "mean"),
        smape_pt_max=("smape_pt", "max"),
        mae_kWh_mean=("mae_kWh", "mean"),
        mae_kWh_max=("mae_kWh", "max"),
        smape_kWh_mean=("smape_kWh", "mean"),
        smape_kWh_max=("smape_kWh", "max"),
        n_windows=("window", "count")
    ).reset_index()

    summary_df["mae_pt_pm"] = summary_df["mae_pt_mean"].map(lambda x: f"{x:.2f}") + "±" + \
                               (summary_df["mae_pt_max"] - summary_df["mae_pt_mean"]).map(lambda x: f"{x:.2f}")
    summary_df["smape_pt_pm"] = summary_df["smape_pt_mean"].map(lambda x: f"{x:.2f}") + "±" + \
                                 (summary_df["smape_pt_max"] - summary_df["smape_pt_mean"]).map(lambda x: f"{x:.2f}")
    summary_df["mae_kWh_pm"] = summary_df["mae_kWh_mean"].map(lambda x: f"{x:.2f}") + "±" + \
                                (summary_df["mae_kWh_max"] - summary_df["mae_kWh_mean"]).map(lambda x: f"{x:.2f}")
    summary_df["smape_kWh_pm"] = summary_df["smape_kWh_mean"].map(lambda x: f"{x:.2f}") + "±" + \
                                  (summary_df["smape_kWh_max"] - summary_df["smape_kWh_mean"]).map(lambda x: f"{x:.2f}")

    summary_path = os.path.join(base_cfg["RESULTS_ROOT"], "sensitivity_summary.csv")
    summary_df.to_csv(summary_path, index=False, encoding="utf-8-sig")

    if not summary_df.empty:
        best_pt_idx = summary_df.groupby("param_name")["mae_pt_mean"].idxmin()
        best_kwh_idx = summary_df.groupby("param_name")["mae_kWh_mean"].idxmin()

        best_pt_df = summary_df.loc[best_pt_idx].sort_values("param_name")
        best_kwh_df = summary_df.loc[best_kwh_idx].sort_values("param_name")

        best_pt_df.to_csv(os.path.join(base_cfg["RESULTS_ROOT"], "best_values_for_parking_time.csv"), index=False, encoding="utf-8-sig")
        best_kwh_df.to_csv(os.path.join(base_cfg["RESULTS_ROOT"], "best_values_for_kWhDelivered.csv"), index=False, encoding="utf-8-sig")

    plot_dir = os.path.join(base_cfg["RESULTS_ROOT"], "plots")
    save_sensitivity_plots(summary_df, plot_dir)

    print("\nAll experiments complete.")
    print(f"Formatted raw metrics saved to: {metrics_path}")
    print(f"Formatted raw metrics xlsx saved to: {metrics_xlsx_path}")
    print(f"Numeric raw metrics with baseline saved to: {numeric_with_baseline_path}")
    print(f"Actual-run numeric metrics saved to: {numeric_run_path}")
    print(f"Dynamic trace saved to: {trace_path}")
    print(f"Summary saved to: {summary_path}")
    print(f"Plots saved to: {plot_dir}")


if __name__ == "__main__":
    main()